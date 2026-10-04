"""Reference lists: goods and services whose VAT treatment a document sets.

The zero-rated and exempt goods are listed in the schedules of the VAT (General)
Regulations (SI 273 of 2003) and their amendments. An administrator loads those
lists as a spreadsheet, one row per item, and each transaction description is
matched against the items' keywords. A match classifies the transaction and cites
the document, so "2 packets of sugar" is classified by whatever the loaded
schedule says about sugar.

Matching:
- Keywords are whole words or phrases, case-insensitive. Plurals match too
  ("packet" matches "packets"), and quantities and units in the description are
  ignored because only the keywords are searched for.
- An item is ruled out if any of its exclude words appear ("milk", excluding
  "chocolate", does not match "milk chocolate").
- Only items in force on the transaction date are used.
- When several items match, the one with the longest matching keyword wins
  (the most specific), so "brown sugar" beats "sugar".
"""

import io
import re
from dataclasses import dataclass
from datetime import date, datetime

from openpyxl import Workbook
from openpyxl.styles import Font
from sqlalchemy import delete, select
from sqlalchemy.orm import Session

from .classify import parse_category
from .ingest.parse import parse_date
from .ingest.readers import IngestError, read_table
from .models import VAT_CATEGORIES, ReferenceItem

COLUMNS = ["Item", "Keywords", "VAT category", "Exclude words", "Source document", "Reference", "From", "To", "Note"]
HEADER_ALIASES = {
    "item": "item", "goods": "item", "description": "item", "name": "item",
    "keywords": "keywords", "keyword": "keywords", "aliases": "keywords", "match": "keywords",
    "vatcategory": "category", "category": "category", "treatment": "category", "vat": "category",
    "excludewords": "excludes", "exclude": "excludes", "excludes": "excludes", "except": "excludes",
    "sourcedocument": "source", "source": "source", "document": "source", "schedule": "source",
    "reference": "reference", "paragraph": "reference", "ref": "reference", "section": "reference",
    "from": "from", "effectivefrom": "from", "start": "from",
    "to": "to", "effectiveto": "to", "end": "to", "until": "to",
    "note": "note", "notes": "note", "comment": "note",
}


def split_terms(text: str | None) -> list[str]:
    return [t.strip().lower() for t in re.split(r"[,;\n]", text or "") if t.strip()]


def _term_regex(term: str) -> re.Pattern:
    words = [re.escape(w) for w in term.split()]
    # Allow a plural on the last word: packet/packets, potato/potatoes, box/boxes.
    return re.compile(r"\b" + r"\s+".join(words) + r"(?:s|es)?\b", re.IGNORECASE)


@dataclass
class ItemMatch:
    item: ReferenceItem
    keyword: str


class ReferenceMatcher:
    def __init__(self, items: list[ReferenceItem]):
        self.entries = []
        for item in items:
            if not item.active:
                continue
            keywords = split_terms(item.keywords) or split_terms(item.item)
            self.entries.append((item, [(k, _term_regex(k)) for k in keywords],
                                 [_term_regex(x) for x in split_terms(item.excludes)]))

    def match(self, text: str, on: date | None = None) -> ItemMatch | None:
        best: ItemMatch | None = None
        for item, keywords, excludes in self.entries:
            if not item.applies_on(on) or any(x.search(text) for x in excludes):
                continue
            for keyword, rx in keywords:
                if rx.search(text) and (best is None or len(keyword) > len(best.keyword)):
                    best = ItemMatch(item, keyword)
        return best


def load_matcher(session: Session) -> ReferenceMatcher:
    return ReferenceMatcher(list(session.scalars(select(ReferenceItem).where(ReferenceItem.active.is_(True)))))


@dataclass
class ImportResult:
    imported: int
    errors: list[str]
    replaced: int = 0


def _map_header(row: list) -> dict[str, int]:
    mapping = {}
    for i, cell in enumerate(row):
        key = HEADER_ALIASES.get(re.sub(r"[^a-z]", "", str(cell or "").lower()))
        if key and key not in mapping:
            mapping[key] = i
    return mapping


def _cell_date(value) -> date | None:
    if value in (None, ""):
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    return parse_date(value)


def import_items(session: Session, filename: str, content: bytes, default_source: str = "",
                 replace_source: bool = True) -> ImportResult:
    """Load items from a CSV or Excel list. With `replace_source`, items previously loaded from the
    same source document are replaced, so an updated schedule can be re-imported."""
    rows = read_table(filename, content)
    header_idx = next((i for i, r in enumerate(rows[:20]) if {"item", "category"} <= _map_header(r).keys()), None)
    if header_idx is None:
        raise IngestError("Couldn't find the heading row. The list needs at least 'Item' and 'VAT category' "
                          "columns; download the template to see the layout.")
    cols = _map_header(rows[header_idx])

    def get(row, key):
        i = cols.get(key)
        value = row[i] if i is not None and i < len(row) else None
        return value.strip() if isinstance(value, str) else value

    items, errors = [], []
    for n, row in enumerate(rows[header_idx + 1:], start=header_idx + 2):
        name = str(get(row, "item") or "").strip()
        if not name and not get(row, "keywords"):
            continue
        category = parse_category(str(get(row, "category") or ""))
        source = str(get(row, "source") or default_source).strip()
        problems = []
        if not name:
            problems.append("no item name")
        if not category:
            problems.append(f"VAT category '{get(row, 'category') or ''}' not recognised "
                            f"(use {', '.join(VAT_CATEGORIES.values())})")
        if not source:
            problems.append("no source document (fill the column or enter one on the import form)")
        start, end = _cell_date(get(row, "from")), _cell_date(get(row, "to"))
        if get(row, "from") not in (None, "") and start is None or get(row, "to") not in (None, "") and end is None:
            problems.append("dates must be like 2026-01-01")
        if problems:
            errors.append(f"Row {n}: " + "; ".join(problems))
            continue
        keywords = str(get(row, "keywords") or name)
        items.append(ReferenceItem(
            item=name, keywords=", ".join(split_terms(keywords)), excludes=", ".join(split_terms(
                str(get(row, "excludes") or ""))) or None, vat_category=category, source=source,
            reference=str(get(row, "reference") or "").strip() or None, effective_from=start, effective_to=end,
            note=str(get(row, "note") or "").strip() or None))
    if errors:
        return ImportResult(0, errors)
    replaced = 0
    if replace_source:
        for source in {i.source for i in items}:
            replaced += session.execute(delete(ReferenceItem).where(ReferenceItem.source == source)).rowcount
    session.add_all(items)
    return ImportResult(len(items), [], replaced)


def template_xlsx() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Reference list"
    # Instructions sit above the heading row; the importer finds the heading row itself.
    for line in [
        "One row per good or service listed in the document (for example each item in a schedule of SI 273 of 2003).",
        "Item: the name as the document gives it.",
        "Keywords: words or phrases that appear in transaction descriptions for this item, separated by commas. "
        "Plurals match automatically. If empty, the item name is used.",
        "VAT category: Zero-rated, Exempt, Standard-rated or Out of scope.",
        "Exclude words: if any of these appear, the item does not match (e.g. item 'Milk', exclude 'chocolate, flavoured').",
        "Source document: e.g. 'SI 273 of 2003, Second Schedule'. Re-importing a list replaces the items with the same source.",
        "Reference: paragraph or item number in the document.",
        "From / To: dates the treatment applies (YYYY-MM-DD). Leave empty if it always applies.",
        "Copy the wording from the document itself; do not guess an item's treatment.",
    ]:
        ws.append([line])
        ws.cell(row=ws.max_row, column=1).font = Font(italic=True, color="666666")
    ws.append([])
    ws.append(COLUMNS)
    for cell in ws[ws.max_row]:
        cell.font = Font(bold=True)
    ws.freeze_panes = ws.cell(row=ws.max_row + 1, column=1)
    for col, width in zip("ABCDEFGHI", (28, 40, 14, 24, 36, 14, 12, 12, 40)):
        ws.column_dimensions[col].width = width
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
