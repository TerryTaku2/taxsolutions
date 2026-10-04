"""Read an uploaded file into a list of rows (lists of cell values)."""

import csv
import io
from pathlib import PurePath


class IngestError(Exception):
    pass


SUPPORTED_EXTENSIONS = {".csv", ".txt", ".xlsx", ".xlsm"}


def read_table(filename: str, content: bytes) -> list[list]:
    ext = PurePath(filename).suffix.lower()
    if ext in (".xlsx", ".xlsm"):
        rows = _read_excel(content)
    elif ext in (".csv", ".txt"):
        rows = _read_csv(content)
    elif ext == ".xls":
        raise IngestError("Old .xls files are not supported. Open the file in Excel and save it as .xlsx or .csv.")
    else:
        raise IngestError(f"Unsupported file type '{ext}'. Upload a CSV or Excel (.xlsx) file.")
    return _trim(rows)


def _read_excel(content: bytes) -> list[list]:
    from openpyxl import load_workbook

    try:
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as e:  # openpyxl raises many types for corrupt files
        raise IngestError(f"Could not open the Excel file: {e}") from e
    best: list[list] = []
    # Use the sheet with the most non-empty rows; exports often carry a cover sheet.
    for ws in wb.worksheets:
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
        filled = [r for r in rows if any(c not in (None, "") for c in r)]
        if len(filled) > len([r for r in best if any(c not in (None, "") for c in r)]):
            best = rows
    wb.close()
    return best


def _read_csv(content: bytes) -> list[list]:
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            text = content.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    else:
        raise IngestError("Could not read the file's text encoding. Save it as UTF-8 CSV.")
    sample = text[:20000]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    return [row for row in csv.reader(io.StringIO(text), dialect)]


def _trim(rows: list[list]) -> list[list]:
    cleaned = []
    for row in rows:
        cells = [c.strip() if isinstance(c, str) else c for c in row]
        while cells and cells[-1] in (None, ""):
            cells.pop()
        cleaned.append(cells)
    while cleaned and not cleaned[-1]:
        cleaned.pop()
    return cleaned
