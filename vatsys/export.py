"""Return-ready schedules as Excel or CSV."""

import csv
import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .checks import Issue
from .compute import ReturnResult
from .models import INPUT_TYPES, PURCHASE, SALE, VAT_CATEGORIES, Business

RETURN_LINES = ["SR_SUPPLIES", "ZR_SUPPLIES", "EX_SUPPLIES", "TOTAL_SUPPLIES", "OUTPUT_TAX",
                "IMPORTED_SERVICES_TAX", "TOTAL_OUTPUT", "INPUT_CAPITAL", "INPUT_IMPORTS", "INPUT_OTHER",
                "TOTAL_INPUT", "NET_VAT"]
INFO_LINES = ["PURCHASES_STANDARD", "PURCHASES_NO_VAT", "NON_DEDUCTIBLE", "OUT_OF_SCOPE"]

BOLD = Font(bold=True)
HEAD_FILL = PatternFill("solid", fgColor="E8EEF4")
MONEY = "#,##0.00;(#,##0.00)"


def _header(ws, row: int, values: list[str]) -> None:
    for col, v in enumerate(values, start=1):
        cell = ws.cell(row=row, column=col, value=v)
        cell.font, cell.fill = BOLD, HEAD_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="top")


def _autosize(ws, widths: dict[int, int] | None = None) -> None:
    for col_cells in ws.columns:
        col = col_cells[0].column
        width = max((len(str(c.value)) for c in col_cells if c.value is not None), default=8)
        ws.column_dimensions[get_column_letter(col)].width = min(max(10, width + 2), 50)
    for col, w in (widths or {}).items():
        ws.column_dimensions[get_column_letter(col)].width = w


def return_to_xlsx(result: ReturnResult, business: Business, issues: list[Issue]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "VAT Return"
    p = result.period
    info = [("Business", business.name), ("BP number", business.bp_number or ""),
            ("VAT number", business.vat_number or ""), ("Tax period", f"{p.start:%d %b %Y} to {p.end:%d %b %Y}"),
            ("Due date", f"{p.due_date:%d %b %Y}"), ("Reporting currency", result.reporting_currency)]
    for r, (k, v) in enumerate(info, start=1):
        ws.cell(row=r, column=1, value=k).font = BOLD
        ws.cell(row=r, column=2, value=v)
    if result.taxable_share is not None:
        info.append(("Taxable share of supplies", f"{result.taxable_share * 100:.1f}% (mixed-use input tax "
                     f"claimed at {result.claim_ratio * 100:.1f}%)"))
    for ccy, net in result.net_by_currency().items():
        info.append((f"Net VAT in {ccy}", f"{net:,.2f} {'payable' if net > 0 else 'refundable'} in {ccy}"))
    for r, (k, v) in enumerate(info[6:], start=7):
        ws.cell(row=r, column=1, value=k).font = BOLD
        ws.cell(row=r, column=2, value=v)
    row = len(info) + 2
    ccys = result.currencies
    _header(ws, row, ["Line", "Description"] + [f"{c} (transaction currency)" for c in ccys]
            + [f"Total in {result.reporting_currency}", "Transactions"])
    for codes in (RETURN_LINES, INFO_LINES):
        for code in codes:
            row += 1
            line = result.line(code)
            values = [code, line.label] + [line.by_currency.get(c) for c in ccys] + [line.total, len(line.txn_ids)]
            for col, v in enumerate(values, start=1):
                cell = ws.cell(row=row, column=col, value=v)
                if 3 <= col <= 3 + len(ccys):
                    cell.number_format = MONEY
                if code in ("TOTAL_SUPPLIES", "TOTAL_OUTPUT", "TOTAL_INPUT", "NET_VAT"):
                    cell.font = BOLD
        row += 1
    _autosize(ws, {2: 45})

    cols = ["Txn ID", "Date", "Invoice no", "Reference", "Counterparty", "Counterparty VAT no", "Description",
            "VAT category", "Input type", "Currency", "Amount", "VAT rate %", "Net", "VAT", "FX rate", "FX rate date",
            f"Net ({result.reporting_currency})", f"VAT ({result.reporting_currency})", "Source file", "Row",
            "VAT claimed", "Claim share"]
    for title, direction in (("Sales schedule", SALE), ("Purchases schedule", PURCHASE)):
        sheet = wb.create_sheet(title)
        _header(sheet, 1, cols)
        r = 1
        for c in result.calcs.values():
            t = c.txn
            if t.direction != direction:
                continue
            r += 1
            values = [t.id, t.txn_date, t.invoice_number, t.reference, t.counterparty, t.counterparty_vat_number,
                      t.description, VAT_CATEGORIES[t.vat_category],
                      INPUT_TYPES.get(t.input_type, "") if direction == PURCHASE else "", t.currency, t.amount,
                      c.rate, c.net, c.vat, c.fx_rate, c.fx_date, c.net_rep, c.vat_rep,
                      t.upload.filename if t.upload else "manual", t.row_number,
                      c.claim if direction == PURCHASE else None,
                      float(c.claim_ratio) if direction == PURCHASE and t.vat_category == "standard" else None]
            if c.error:
                values[12:18] = ["ERROR: " + c.error, None, None, None, None, None]
            for col, v in enumerate(values, start=1):
                cell = sheet.cell(row=r, column=col, value=v)
                if col in (11, 13, 14, 17, 18, 21):
                    cell.number_format = MONEY
                elif col == 22:
                    cell.number_format = "0.0%"
                elif col in (2, 16):
                    cell.number_format = "dd/mm/yyyy"
        _autosize(sheet)
        sheet.freeze_panes = "A2"

    fx = wb.create_sheet("Exchange rates")
    _header(fx, 1, ["From", "To", "Rate date", "Rate", "Source", "Transactions converted"])
    for r, f in enumerate(result.fx_used, start=2):
        for col, v in enumerate([f["from"], f["to"], f["date"], f["rate"], f["source"], f["count"]], start=1):
            cell = fx.cell(row=r, column=col, value=v)
            if col == 3:
                cell.number_format = "dd/mm/yyyy"
    _autosize(fx)

    chk = wb.create_sheet("Checks")
    _header(chk, 1, ["Severity", "Check", "Detail", "Transactions"])
    for r, i in enumerate(issues, start=2):
        for col, v in enumerate([i.severity, i.title, i.detail, ", ".join(map(str, i.txn_ids[:50]))], start=1):
            chk.cell(row=r, column=col, value=v)
    _autosize(chk, {3: 60, 4: 40})

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def return_to_csv(result: ReturnResult) -> str:
    buf = io.StringIO()
    w = csv.writer(buf)
    ccys = result.currencies
    w.writerow(["Line", "Description"] + ccys + [f"Total {result.reporting_currency}"])
    for code in RETURN_LINES + INFO_LINES:
        line = result.line(code)
        w.writerow([code, line.label] + [f"{line.by_currency.get(c):.2f}" for c in ccys] + [f"{line.total:.2f}"])
    return buf.getvalue()
