"""Reference lists: items loaded from documents classify transactions by what they describe.

The items here are test data, not the real schedules of the VAT (General) Regulations.
"""

import io
import re
from datetime import date

from openpyxl import load_workbook
from sqlalchemy import select

from tests.conftest import csrf_from
from tests.test_web import register
from vatsys.ingest import classifier_for
from vatsys.models import ClassificationRule, ReferenceItem
from vatsys.reference import import_items, template_xlsx

LIST_CSV = b"""Item,Keywords,VAT category,Exclude words,Source document,Reference,From,To
Sugar,sugar,Zero-rated,,TEST Schedule A,para 1,,
Brown sugar,brown sugar,Exempt,,TEST Schedule A,para 2,,
Fresh milk,milk,Zero-rated,"chocolate, flavoured",TEST Schedule A,para 3,,
Sunflower oil,cooking oil,Standard-rated,,TEST Schedule B,item 9,2026-01-01,
"""


def load(session, content=LIST_CSV, **kw):
    result = import_items(session, "list.csv", content, **kw)
    session.flush()
    return result


def classify(session, business, text, on=date(2026, 9, 1), direction="sale"):
    return classifier_for(session, business.id).classify(direction=direction, description=text,
                                                         counterparty=None, on=on)


def test_item_in_description_is_classified(session, business):
    assert load(session).imported == 4
    c = classify(session, business, "Sale of 2 packets of sugar")
    assert (c.vat_category, c.source) == ("zero", "reference")
    assert "TEST Schedule A, para 1" in c.reason and c.reference_item_id
    assert c.confidence >= 0.8  # not flagged for review


def test_most_specific_item_wins(session, business):
    load(session)
    assert classify(session, business, "Brown sugar 2kg").vat_category == "exempt"
    assert classify(session, business, "SUGAR 10 x 2kg").vat_category == "zero"


def test_excluded_words_and_dates(session, business):
    load(session)
    assert classify(session, business, "Milk chocolate bars").source == "default"
    assert classify(session, business, "Fresh milk 2L x 3").vat_category == "zero"
    assert classify(session, business, "Cooking oil 2L", on=date(2025, 6, 1)).source == "default"
    assert classify(session, business, "Cooking oil 2L").vat_category == "standard"


def test_business_rule_beats_list_and_list_beats_global_rule(session, business):
    load(session)
    session.add(ClassificationRule(business_id=business.id, pattern="sugar", match_type="contains", field="any",
                                   direction="any", vat_category="exempt", priority=100))
    session.add(ClassificationRule(business_id=None, pattern="milk", match_type="contains", field="any",
                                   direction="any", vat_category="exempt", priority=1))
    session.flush()
    assert classify(session, business, "sugar 2kg").source == "rule"
    assert classify(session, business, "Fresh milk").source == "reference"


def test_reimport_replaces_same_source(session):
    load(session)
    r = load(session, b"Item,VAT category,Source document\nSalt,Zero-rated,TEST Schedule A\n")
    assert (r.imported, r.replaced) == (1, 3)
    assert {i.item for i in session.scalars(select(ReferenceItem))} == {"Salt", "Sunflower oil"}


def test_invalid_rows_import_nothing(session):
    r = load(session, b"Item,VAT category,Source document\nSalt,Maybe,TEST\nRice,Zero-rated,\n")
    assert r.imported == 0 and len(r.errors) == 2
    assert "not recognised" in r.errors[0] and "no source" in r.errors[1]
    assert load(session, b"Item,VAT category\nRice,Zero-rated\n", default_source="TEST C").imported == 1


def test_template_has_expected_columns(session):
    wb = load_workbook(io.BytesIO(template_xlsx()))
    rows = [list(r) for r in wb.active.iter_rows(values_only=True)]
    assert rows[-1][:3] == ["Item", "Keywords", "VAT category"]
    result = import_items(session, "t.xlsx", template_xlsx())
    assert result.imported == 0 and not result.errors


def test_web_import_and_upload_uses_list(client):
    token = register(client)
    r = client.post("/settings/reference/import", data={"csrf": token, "replace_source": "yes"},
                    files={"file": ("list.csv", LIST_CSV)})
    assert "Imported 4 item(s)" in r.text
    r = client.get("/settings/reference", params={"test": "2 packets of sugar", "test_date": "2026-09-01"})
    assert "Zero-rated" in r.text and "TEST Schedule A, para 1" in r.text

    r = client.post("/businesses/new", data={"csrf": token, "name": "Shop"})
    bid = int(re.search(r"/b/(\d+)", str(r.url)).group(1))
    sales = b"Date,Description,Amount\n2026-09-02,2 packets of sugar,4.00\n2026-09-02,Bread,1.00\n"
    client.post(f"/b/{bid}/uploads", data={"csrf": token, "source_type": "ledger", "default_direction": "sale",
                                           "default_currency": "USD", "amounts_include_vat": "yes"},
                files={"file": ("sales.csv", sales)})
    r = client.get(f"/b/{bid}/transactions")
    assert ">List<" in r.text and ">Check<" in r.text  # sugar from the list; bread not listed


def test_only_admin_imports(client):
    token = register(client)
    client.post("/businesses/new", data={"csrf": token, "name": "A", "owner_email": "o@example.com",
                                         "temp_password": "temporary-pass-1"})
    client.cookies.clear()
    token = csrf_from(client.get("/login").text)
    r = client.post("/login", data={"csrf": token, "email": "o@example.com", "password": "temporary-pass-1"})
    token = csrf_from(r.text)
    client.post("/account/password", data={"csrf": token, "current_password": "temporary-pass-1",
                                           "new_password": "owner-password-1", "confirm_password": "owner-password-1"})
    r = client.post("/settings/reference/import", data={"csrf": token}, files={"file": ("list.csv", LIST_CSV)})
    assert r.status_code == 403
